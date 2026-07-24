#!/usr/bin/env python3
"""Normalize a Numbers-exported parliamentary election workbook.

Usage:
    python3 scripts/extract_data.py /path/to/pru15.xlsx public/data/elections/pru-15/election.json
"""

from __future__ import annotations

import json
import argparse
import re
import unicodedata
from collections import Counter, defaultdict
from pathlib import Path

from openpyxl import load_workbook


ALLIANCES = {
    "PAKATAN HARAPAN (PH)": ("PH", "#dc4b43"),
    "PERIKATAN NASIONAL (PN)": ("PN", "#186a58"),
    "BARISAN NASIONAL (BN)": ("BN", "#245c9f"),
    "GABUNGAN PARTI SARAWAK (GPS)": ("GPS", "#e1a72b"),
    "PARTI GABUNGAN RAKYAT SABAH (GRS)": ("GRS", "#65a7ba"),
    "PARTI PEJUANG TANAHAIR (PEJUANG)": ("PEJUANG", "#8b744f"),
    "PARTI BUMIPUTERA PERKASA MALAYSIA (PUTRA)": ("PUTRA", "#8c5899"),
    "LAIN-LAIN / BEBAS": ("LAIN-LAIN", "#7b8580"),
}

# Source workbook rows disagree for this seat. The SPR age-by-DUN report and
# the confirmed constituency total both record 43,072 registered voters.
REGISTERED_VOTER_CORRECTIONS = {
    "pru-15": {"P.201": 43_072},
}


def clean(value):
    return value.strip() if isinstance(value, str) else value


def alliance_name(value):
    value = clean(value)
    return value if value else "LAIN-LAIN / BEBAS"


def round_fraction(value):
    return round(float(value or 0), 6)


def candidate_id(seat_code: str, name: str) -> str:
    ascii_name = unicodedata.normalize("NFKD", name).encode("ascii", "ignore").decode("ascii")
    slug = re.sub(r"[^a-z0-9]+", "-", ascii_name.lower()).strip("-")
    return f"{seat_code.lower().replace('.', '')}:{slug}"


def normalise_name(value: str) -> str:
    ascii_name = unicodedata.normalize("NFKD", value).encode("ascii", "ignore").decode("ascii")
    return re.sub(r"[^a-z0-9]+", " ", ascii_name.lower()).strip()


def person_lookups(persons_path: Path, matches_path: Path) -> tuple[dict[str, set[str]], dict[tuple[str, str, str], str]]:
    by_name: dict[str, set[str]] = defaultdict(set)
    if persons_path.exists():
        registry = json.loads(persons_path.read_text(encoding="utf-8"))
        for person in registry.get("persons", []):
            for name in [person.get("canonicalName", ""), *person.get("aliases", [])]:
                if name:
                    by_name[normalise_name(name)].add(person["id"])
    explicit: dict[tuple[str, str, str], str] = {}
    if matches_path.exists():
        matches = json.loads(matches_path.read_text(encoding="utf-8"))
        for match in matches.get("matches", []):
            explicit[(match["electionId"], match["seatCode"], normalise_name(match["candidateName"]))] = match["personId"]
    return by_name, explicit


def main() -> None:
    parser = argparse.ArgumentParser(description="Normalize a parliamentary election workbook.")
    parser.add_argument("input", type=Path)
    parser.add_argument("output", type=Path)
    parser.add_argument("--election-number", type=int, default=15)
    parser.add_argument("--election-date", default="2022-11-19")
    parser.add_argument("--term-id")
    parser.add_argument("--boundary-version", default="my-sarawak-2015-peninsula-2018-sabah-2019")
    parser.add_argument("--title")
    parser.add_argument("--short-title")
    parser.add_argument("--source-file")
    parser.add_argument("--expected-seats", type=int, default=222)
    parser.add_argument("--expected-candidates", type=int, default=945)
    parser.add_argument("--persons", type=Path, default=Path("public/data/reference/persons.json"))
    parser.add_argument("--person-matches", type=Path, default=Path("public/data/reference/person-matches.json"))
    args = parser.parse_args()

    input_path = args.input.resolve()
    output_path = args.output.resolve()
    election_id = f"pru-{args.election_number}"
    term_id = args.term_id or f"dr-{args.election_number}"
    short_title = args.short_title or f"PRU-{args.election_number}"
    title = args.title or f"Pilihan Raya Umum Malaysia ke-{args.election_number}"
    corrections = REGISTERED_VOTER_CORRECTIONS.get(election_id, {})
    people_by_name, explicit_person_matches = person_lookups(args.persons, args.person_matches)
    workbook = load_workbook(input_path, read_only=True, data_only=True)

    results_sheet = workbook["ALL RESULT - Table 1"]
    winner_sheet = workbook["WINNERS"]

    result_rows = [
        tuple(clean(cell) for cell in row)
        for row in results_sheet.iter_rows(min_row=2, values_only=True)
        if row[1]
    ]
    winner_rows = [
        tuple(clean(cell) for cell in row)
        for row in winner_sheet.iter_rows(min_row=2, values_only=True)
        if row[1]
    ]

    winners = {row[1]: row for row in winner_rows}
    grouped = defaultdict(list)
    for row in result_rows:
        grouped[row[1]].append(row)

    seats = []
    issues = []
    for code in sorted(grouped, key=lambda item: int(item.split(".")[1])):
        rows = grouped[code]
        winner_row = winners.get(code)
        if winner_row is None:
            raise ValueError(f"No winner row for {code}")

        winner_candidate = winner_row[3]
        source_winner = next((row for row in rows if row[3] == winner_candidate), None)
        if source_winner is None:
            source_winner = max(rows, key=lambda row: row[8] or 0)

        registered_values = sorted({int(row[6]) for row in rows if row[6] is not None})
        turnout_values = sorted({int(row[7]) for row in rows if row[7] is not None})
        if (len(registered_values) > 1 or len(turnout_values) > 1) and code not in corrections:
            issues.append(
                {
                    "seat": code,
                    "field": "registeredVoters",
                    "message": (
                        "Rekod calon mempunyai jumlah pemilih berdaftar yang bercanggah "
                        f"({', '.join(map(str, registered_values))}). Nilai pada rekod pemenang digunakan."
                    ),
                }
            )

        candidates = sorted(
            [
                {
                    "id": candidate_id(code, row[3]),
                    "name": row[3],
                    "alliance": alliance_name(row[4]),
                    "party": row[5] or "TIDAK DINYATAKAN",
                    "votes": int(row[8] or 0),
                    "share": round_fraction(row[9]),
                }
                for row in rows
            ],
            key=lambda item: item["votes"],
            reverse=True,
        )

        runner_up = candidates[1] if len(candidates) > 1 else None
        winner = candidates[0]
        registered = corrections.get(code, int(source_winner[6] or 0))
        turnout = int(source_winner[7] or 0)
        seats.append(
            {
                "code": code,
                "state": rows[0][0],
                "name": rows[0][2],
                "registered": registered,
                "turnout": turnout,
                "turnoutPct": round_fraction(turnout / registered if registered else 0),
                "candidateCount": len(candidates),
                "marginVotes": winner["votes"] - (runner_up["votes"] if runner_up else 0),
                "marginShare": round_fraction(
                    winner["share"] - (runner_up["share"] if runner_up else 0)
                ),
                "winner": {
                    **winner,
                    "gender": winner_row[4] or "TIDAK DINYATAKAN",
                    "ethnicity": winner_row[5] or "TIDAK DINYATAKAN",
                },
                "candidates": candidates,
            }
        )

    if len(seats) != args.expected_seats:
        raise ValueError(f"Expected {args.expected_seats} seats, found {len(seats)}")
    if args.expected_candidates and sum(len(seat["candidates"]) for seat in seats) != args.expected_candidates:
        raise ValueError(f"Expected {args.expected_candidates} candidates; source workbook does not reconcile")

    person_slugs = [candidate["id"].split(":", 1)[1] for seat in seats for candidate in seat["candidates"]]
    duplicate_people = {slug for slug, count in Counter(person_slugs).items() if count > 1}
    for seat in seats:
        seat["electionId"] = election_id
        seat["contestId"] = f"{election_id}:{seat['code']}"
        for index, candidate in enumerate(seat["candidates"], start=1):
            person_slug = candidate["id"].split(":", 1)[1]
            suffix = f":{seat['code'].lower().replace('.', '')}" if person_slug in duplicate_people else ""
            explicit_person_id = explicit_person_matches.get((election_id, seat["code"], normalise_name(candidate["name"])))
            existing_person_ids = people_by_name.get(normalise_name(candidate["name"]), set())
            person_id = explicit_person_id or (next(iter(existing_person_ids)) if len(existing_person_ids) == 1 else f"person:{person_slug}{suffix}")
            candidate.update(
                {
                    "electionId": election_id,
                    "candidacyId": f"{election_id}:{seat['code']}:{index:02d}",
                    "personId": person_id,
                }
            )
        winner_identity = next(candidate for candidate in seat["candidates"] if candidate["id"] == seat["winner"]["id"])
        seat["winner"].update(
            {
                "electionId": winner_identity["electionId"],
                "candidacyId": winner_identity["candidacyId"],
                "personId": winner_identity["personId"],
            }
        )

    alliances = [
        {"name": name, "shortName": short_name, "color": color}
        for name, (short_name, color) in ALLIANCES.items()
    ]
    payload = {
        "metadata": {
            "electionId": election_id,
            "electionNumber": args.election_number,
            "termId": term_id,
            "boundaryVersion": args.boundary_version,
            "title": title,
            "shortTitle": short_title,
            "electionDate": args.election_date,
            "sourceFile": args.source_file or ("DATA & STATISTIK PRU-15.numbers" if args.election_number == 15 else input_path.name),
            "seatCount": len(seats),
            "candidateCount": sum(len(seat["candidates"]) for seat in seats),
            "stateCount": len({seat["state"] for seat in seats}),
            "issues": issues,
        },
        "alliances": alliances,
        "seats": seats,
    }

    output_path.parent.mkdir(parents=True, exist_ok=True)
    output_path.write_text(
        json.dumps(payload, ensure_ascii=False, separators=(",", ":")), encoding="utf-8"
    )
    print(
        f"Wrote {len(seats)} seats / {payload['metadata']['candidateCount']} candidates "
        f"to {output_path}"
    )


if __name__ == "__main__":
    main()
